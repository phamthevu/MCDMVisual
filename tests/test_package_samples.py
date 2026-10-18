import unittest
from pathlib import Path
import numpy as np
from app.core.dematel import compute_dematel_from_total
from app.core.io_excel import load_total_influence_matrix, load_ranking_results

ROOT = Path(__file__).resolve().parents[1]

class PackagedSampleTests(unittest.TestCase):
    def test_dematel_demo_is_valid_and_weights_normalize(self):
        labels, matrix = load_total_influence_matrix(str(ROOT / 'Demo Data' / 'DEMATEL_Demo.xlsx'))
        rows = compute_dematel_from_total(matrix, labels)
        weights = np.array([row['Weight'] for row in rows])
        self.assertTrue(np.isfinite(matrix).all())
        self.assertAlmostEqual(float(weights.sum()), 1.0, places=12)
        self.assertTrue(all('R+C' in row and 'R-C' in row for row in rows))

    def test_ranking_demo_and_template_layout_are_readable(self):
        labels, cols, values = load_ranking_results(str(ROOT / 'Demo Data' / 'Ranking_Demo.xlsx'))
        self.assertEqual(len(labels), 4)
        self.assertEqual(cols, ['Score', 'Rank'])
        self.assertEqual(values.shape, (4, 2))
        # Check the blank workbook preserves the same input schema.
        from openpyxl import load_workbook
        sheet = load_workbook(ROOT / 'Template' / 'Ranking_Results_Template.xlsx', read_only=True).active
        self.assertEqual([c.value for c in sheet[1]], ['Alternative', 'Score', 'Rank'])
        self.assertGreater(sheet.max_row, 1)

if __name__ == '__main__':
    unittest.main()
